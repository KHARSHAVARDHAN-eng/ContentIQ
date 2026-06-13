# ==========================================
# PRODUCTION LOCKED - STABLE RAG V1 CORE
# DO NOT MODIFY without explicit regression verification
# ==========================================

import os
import traceback
import pdfplumber
import docx
from sqlalchemy.orm import Session
from app.core.database import SessionLocal
from app.models.document import Document
from app.models.document_page import DocumentPage
from app.models.document_chunk import DocumentChunk
from app.models.chunk_embedding import ChunkEmbedding
from app.core.config import settings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.services.embedding_service import embedding_service

def process_document_task(document_id: int):
    db = SessionLocal()
    try:
        doc = db.query(Document).filter(Document.id == document_id).first()
        if not doc:
            print(f"Document {document_id} not found in database.")
            return

        print(f"Starting extraction for document {document_id}: {doc.name}")
        doc.status = "PROCESSING"
        db.commit()

        # Clean up existing pages, chunks, and embeddings to prevent duplicates on rerun
        db.query(ChunkEmbedding).filter(
            ChunkEmbedding.chunk_id.in_(
                db.query(DocumentChunk.id).filter(DocumentChunk.document_id == document_id)
            )
        ).delete(synchronize_session=False)
        db.query(DocumentChunk).filter(DocumentChunk.document_id == document_id).delete(synchronize_session=False)
        db.query(DocumentPage).filter(DocumentPage.document_id == document_id).delete(synchronize_session=False)
        db.commit()

        # Use storage_service to get file bytes
        try:
            import io
            from app.services.storage_service import storage_service
            file_bytes = storage_service.download_file(doc.path)
        except Exception as e:
            print(f"Error downloading file from storage: {e}")
            doc.status = "FAILED"
            db.commit()
            return

        _, ext = os.path.splitext(doc.name.lower())
        pages_to_insert = []
        
        if ext == ".pdf":
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for idx, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    text_content = text.strip() if text else "[Empty Page]"
                    if not text_content:
                        text_content = "[Empty Page]"
                    
                    db_page = DocumentPage(
                        document_id=doc.id,
                        page_number=idx + 1,
                        extracted_text=text_content
                    )
                    pages_to_insert.append(db_page)
                    
        elif ext == ".docx":
            word_doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in word_doc.paragraphs:
                full_text.append(para.text)
            
            for table in word_doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    full_text.append(" | ".join(row_text))

            text_content = "\n".join(full_text).strip()
            if not text_content:
                text_content = "[Empty Document]"

            db_page = DocumentPage(
                document_id=doc.id,
                page_number=1,
                extracted_text=text_content
            )
            pages_to_insert.append(db_page)
            
        else:
            # Plain text files
            try:
                text_content = file_bytes.decode("utf-8").strip()
            except UnicodeDecodeError:
                text_content = file_bytes.decode("latin-1").strip()

            if not text_content:
                text_content = "[Empty Document]"

            db_page = DocumentPage(
                document_id=doc.id,
                page_number=1,
                extracted_text=text_content
            )
            pages_to_insert.append(db_page)

        # Stage 1 Save Pages & transition to TEXT_EXTRACTED
        if pages_to_insert:
            db.bulk_save_objects(pages_to_insert)
            doc.status = "TEXT_EXTRACTED"
            db.commit()
            print(f"Finished extraction. Transitioned document {document_id} to TEXT_EXTRACTED")

            # Transition to CHUNKED state
            doc.status = "CHUNKED"
            db.commit()
            print(f"Starting chunking for document {document_id}...")

            text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=settings.CHUNK_SIZE,
                chunk_overlap=settings.CHUNK_OVERLAP,
                length_function=len
            )

            chunks_to_insert = []
            chunk_global_index = 0

            for page in pages_to_insert:
                text = page.extracted_text
                if text in ["[Empty Page]", "[Empty Document]"]:
                    continue
                
                splits = text_splitter.split_text(text)
                for split_text in splits:
                    clean_text = split_text.strip()
                    if not clean_text:
                        continue
                    db_chunk = DocumentChunk(
                        document_id=doc.id,
                        page_number=page.page_number,
                        chunk_index=chunk_global_index,
                        chunk_text=clean_text,
                        chunk_length=len(clean_text)
                    )
                    chunks_to_insert.append(db_chunk)
                    chunk_global_index += 1

            if chunks_to_insert:
                # Add chunks to session and flush to obtain their IDs
                db.add_all(chunks_to_insert)
                db.flush()
                
                # Transition status to EMBEDDING_GENERATION
                doc.status = "EMBEDDING_GENERATION"
                db.commit()
                print(f"Document {document_id} chunks created. Starting embedding generation...")

                chunk_texts = [c.chunk_text for c in chunks_to_insert]
                vectors = embedding_service.get_embeddings(chunk_texts)

                embeddings_to_insert = []
                for chunk, vector in zip(chunks_to_insert, vectors):
                    db_emb = ChunkEmbedding(
                        chunk_id=chunk.id,
                        embedding_dimension=len(vector)
                    )
                    embeddings_to_insert.append(db_emb)

                if embeddings_to_insert:
                    db.add_all(embeddings_to_insert)
                
                doc.status = "EMBEDDED"
                db.commit()
                print(f"Finished processing and embedding successfully. Document {document_id} status: EMBEDDED")

                # Transition status to INDEXING
                doc.status = "INDEXING"
                db.commit()
                print(f"Document {document_id}: Starting Qdrant indexing...")

                from app.services.vector_store import vector_store
                vector_store.create_collection()
                
                # Prevent duplication if indexing is rerun
                vector_store.delete_document_vectors(doc.id)

                points_data = []
                for chunk, vector in zip(chunks_to_insert, vectors):
                    points_data.append({
                        "chunk_id": chunk.id,
                        "document_id": doc.id,
                        "page_number": chunk.page_number,
                        "chunk_text": chunk.chunk_text,
                        "vector": vector
                    })

                if points_data:
                    vector_store.upsert_chunks_bulk(points_data)

                doc.status = "INDEXED"
                db.commit()
                print(f"Finished processing, embedding and indexing successfully. Document {document_id} status: INDEXED")
            else:
                doc.status = "INDEXED"
                db.commit()
                print(f"Document {document_id} has no chunks. Transitioned to INDEXED")
        else:
            doc.status = "FAILED"
            db.commit()

        
    except Exception as e:
        print(f"Error during document extraction/chunking/embedding {document_id}: {e}")
        traceback.print_exc()
        try:
            doc = db.query(Document).filter(Document.id == document_id).first()
            if doc:
                doc.status = "FAILED"
                db.commit()
        except Exception as db_err:
            print(f"Failed to set document status to FAILED: {db_err}")
    finally:
        db.close()
